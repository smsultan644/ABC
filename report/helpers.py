"""Word styling helpers."""
from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Inches, Pt, RGBColor, Emu
from docx.enum.style import WD_STYLE_TYPE

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
TEAL = RGBColor(0x1F, 0x6F, 0x6A)
SLATE = RGBColor(0x3D, 0x4A, 0x5C)
RED = RGBColor(0x8B, 0x1E, 0x1E)
CREAM = "F7F4EC"
WHITE = "FFFFFF"
NAVY_HEX = "0B1F3A"
TEAL_HEX = "1F6F6A"
GOLD_HEX = "C4A35A"
ROW_ALT = "F3EFE4"
LIGHT_TEAL = "E6F0EF"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    """kwargs: top, left, bottom, right = {sz, color, val}"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            el = OxmlElement(f"w:{edge}")
            for k, v in kwargs[edge].items():
                el.set(qn(f"w:{k}"), str(v))
            tcBorders.append(el)
    tcPr.append(tcBorders)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def set_narrow_cell_margins(cell, dxa=60):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(dxa))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instr)
    run._r.append(fldChar2)
    set_run_font(run, size=8, color=NAVY)


def add_num_pages(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " NUMPAGES "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instr)
    run._r.append(fldChar2)
    set_run_font(run, size=8, color=NAVY)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].font.color.rgb = SLATE
    styles["Normal"].paragraph_format.space_after = Pt(8)
    styles["Normal"].paragraph_format.space_before = Pt(0)
    styles["Normal"].paragraph_format.line_spacing = 1.12

    for name, size, before, after, color in [
        ("Heading 1", 16, 16, 8, NAVY),
        ("Heading 2", 13, 13, 6, TEAL),
        ("Heading 3", 11.5, 10, 4, NAVY),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    r = hp.add_run("STRATEGIC PASSIVE INCOME & WEALTH-BUILDING BLUEPRINT  ·  2026–2029")
    set_run_font(r, size=8, color=NAVY, bold=True)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    r = fp.add_run("Personal strategic advisory  ·  Not legal, tax or investment advice  ·  Page ")
    set_run_font(r, size=8, color=NAVY)
    add_page_number(fp)
    r = fp.add_run(" of ")
    set_run_font(r, size=8, color=NAVY)
    add_num_pages(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return doc


def p(doc, text, size=11, bold=False, italic=False, color=None, align="left", space_after=8, space_before=0, first_line=0):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.line_spacing = 1.12
    if first_line:
        para.paragraph_format.first_line_indent = Cm(first_line)
    para.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color or SLATE)
    return para


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def bullets(doc, items, size=11):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        para.clear()
        para.paragraph_format.space_after = Pt(3)
        para.paragraph_format.space_before = Pt(0)
        run = para.add_run(it)
        set_run_font(run, size=size, color=SLATE)


def numbered(doc, items, size=11):
    for it in items:
        para = doc.add_paragraph(style="List Number")
        para.clear()
        para.paragraph_format.space_after = Pt(3)
        run = para.add_run(it)
        set_run_font(run, size=size, color=SLATE)


def callout(doc, title, body, fill=LIGHT_TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.text = ""
    p1 = cell.paragraphs[0]
    r = p1.add_run(title)
    set_run_font(r, size=10, bold=True, color=NAVY)
    p1.paragraph_format.space_after = Pt(2)
    p2 = cell.add_paragraph()
    r = p2.add_run(body)
    set_run_font(r, size=10, color=SLATE)
    p2.paragraph_format.space_after = Pt(2)
    set_narrow_cell_margins(cell, 80)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def source_line(doc, text):
    p(doc, text, size=9, italic=True, color=TEAL, space_after=10)


def caption(doc, text):
    p(doc, text, size=9, italic=True, color=NAVY, align="center", space_after=10, space_before=2)


def add_image(doc, path, width_cm=17.2):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run()
    run.add_picture(str(path), width=Cm(width_cm))


def make_table(doc, headers, rows, col_widths=None, header_fill=NAVY_HEX, font_size=8.5):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        run = para.add_run(str(h))
        set_run_font(run, size=font_size, bold=True, color=RGBColor(255, 255, 255))
        shade_cell(cell, header_fill)
        set_narrow_cell_margins(cell, 40)
        prevent_row_split(table.rows[0])
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(str(val))
            set_run_font(run, size=font_size, color=SLATE)
            if r_i % 2 == 1:
                shade_cell(cell, ROW_ALT)
            set_narrow_cell_margins(cell, 40)
        prevent_row_split(table.rows[r_i + 1])
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def page_break(doc):
    doc.add_page_break()


def label_fact(kind):
    return {
        "fact": "Verified fact (public source)",
        "inference": "Reasonable inference",
        "estimate": "Estimate / modelling assumption",
        "opinion": "Strategic opinion",
    }[kind]
